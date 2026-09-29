import requests
import json
import re
import os
from html.parser import HTMLParser
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

BACKEND_URL = 'https://ssc-catalog-cap-backend.cfapps.us10-001.hana.ondemand.com'

# ── HTML stripper ─────────────────────────────────────────────────────────────
class _Stripper(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
    def handle_data(self, data):
        self.parts.append(data)

def strip_html(text):
    if not text:
        return ""
    s = _Stripper()
    s.feed(str(text))
    return re.sub(r'\s+', ' ', ''.join(s.parts)).strip()

# ── Classification helpers ────────────────────────────────────────────────────
def get_classification_values(svc, feature_key):
    cf = svc.get('classificationFeatures')
    if not cf or not isinstance(cf, list):
        return []
    for item in cf:
        if isinstance(item, dict) and item.get('key') == feature_key:
            val = item.get('value', [])
            if isinstance(val, list):
                return [str(v) for v in val]
            return [str(val)] if val else []
    return []

def get_single_classification(svc, feature_key):
    vals = get_classification_values(svc, feature_key)
    return vals[0] if vals else ""

def get_supercategory_keys(svc):
    cats = svc.get('supercategories')
    if not cats or not isinstance(cats, list):
        return []
    result = []
    for c in cats:
        if isinstance(c, dict):
            code = c.get('code') or ''
            name = c.get('name') or ''
            key = f"{code} - {name}" if code and name else code or name
            if key:
                result.append(key)
    return result

def get_sc_keywords(svc):
    kw = svc.get('scKeywords')
    if not kw:
        return ""
    if isinstance(kw, list):
        return '; '.join([k.get('name', str(k)) if isinstance(k, dict) else str(k) for k in kw])
    return str(kw)

# ── Fetch all services with full fields from backend ─────────────────────────
def fetch_all_services():
    print("Fetching all services with full fields from backend...")
    r = requests.get(f"{BACKEND_URL}/api/catalog/export-full", timeout=300)
    r.raise_for_status()
    data = r.json()
    services = data.get('services', [])
    print(f"✅ Loaded {len(services)} services")
    return services

# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    services = fetch_all_services()
    if not services:
        print("❌ No services returned")
        return

    # Debug first service
    s0 = services[0]
    cf = s0.get('classificationFeatures')
    print(f"  CF type: {type(cf)}, length: {len(cf) if isinstance(cf, list) else 'N/A'}")
    if isinstance(cf, list) and cf:
        print(f"  First CF item: {json.dumps(cf[0])[:300]}")

    # ── Collect all unique dynamic column values ──────────────────────────────
    all_et     = set()
    all_phases = set()
    all_sc     = set()

    for svc in services:
        for v in get_classification_values(svc, 'engagementType'):
            all_et.add(v)
        for v in get_classification_values(svc, 'sapActivateProjectPhase'):
            all_phases.add(v)
        for v in get_supercategory_keys(svc):
            all_sc.add(v)

    all_et     = sorted(all_et)
    all_phases = sorted(all_phases)
    all_sc     = sorted(all_sc)

    print(f"  Engagement types ({len(all_et)}): {all_et}")
    print(f"  Project phases   ({len(all_phases)}): {all_phases}")
    print(f"  Supercategories  ({len(all_sc)}): {all_sc[:5]}...")

    # ── Build workbook ────────────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = "Services"

    hdr_font  = Font(bold=True, color="FFFFFF")
    hdr_fill  = PatternFill("solid", fgColor="1F4E79")
    dyn_fill  = PatternFill("solid", fgColor="2E75B6")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    fixed_cols = [
        ("Service Code",      lambda s: s.get('code', '')),
        ("Service Number",    lambda s: s.get('serviceNumber') or s.get('number', '')),
        ("Service Name",      lambda s: s.get('name', '')),
        ("Short Name",        lambda s: s.get('serviceShortName') or s.get('shortName', '')),
        ("CRM Base Category", lambda s: get_single_classification(s, 'crmBaseCategory')),
        ("Efforts (days)",    lambda s: get_single_classification(s, 'effortEstimateDays')),
        ("Summary",           lambda s: strip_html(s.get('summary', ''))),
        ("Teaser Text",       lambda s: strip_html(s.get('serviceTeaserText') or s.get('teaserText', ''))),
        ("Business Needs",    lambda s: strip_html(s.get('businessNeeds', ''))),
        ("Key Benefits",      lambda s: strip_html(s.get('keyBenefits', ''))),
        ("Delivery Approach", lambda s: strip_html(s.get('deliveryApproach', ''))),
        ("Description",       lambda s: strip_html(s.get('description', ''))),
        ("SC Keywords",       lambda s: get_sc_keywords(s)),
    ]

    headers = [c[0] for c in fixed_cols]
    for et in all_et:
        headers.append(f"ET: {et}")
    for ph in all_phases:
        headers.append(f"Phase: {ph}")
    for sc in all_sc:
        headers.append(f"SC: {sc}")

    ws.append(headers)

    for col_idx, cell in enumerate(ws[1], 1):
        cell.font      = hdr_font
        cell.fill      = dyn_fill if col_idx > len(fixed_cols) else hdr_fill
        cell.alignment = hdr_align
        ws.column_dimensions[cell.column_letter].width = 20

    for i in range(7, 13):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 50

    for svc in services:
        row = [fn(svc) for _, fn in fixed_cols]

        svc_et     = set(get_classification_values(svc, 'engagementType'))
        svc_phases = set(get_classification_values(svc, 'sapActivateProjectPhase'))
        svc_sc     = set(get_supercategory_keys(svc))

        for et in all_et:
            row.append("v" if et in svc_et else "")
        for ph in all_phases:
            row.append("v" if ph in svc_phases else "")
        for sc in all_sc:
            row.append("v" if sc in svc_sc else "")

        ws.append(row)

    ws.freeze_panes = "A2"

    output_file = os.path.join(os.path.expanduser("~"), "Desktop", "ssc_services_export.xlsx")
    wb.save(output_file)
    print(f"\n✅ Saved: {output_file}")
    print(f"   {len(services)} services × {len(headers)} columns")

if __name__ == "__main__":
    main()
