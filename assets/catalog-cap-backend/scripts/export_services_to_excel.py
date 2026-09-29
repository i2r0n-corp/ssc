#!/usr/bin/env python3
"""
Export all services from SSC Catalog data product to Excel.
"""

import requests
import json
import re
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BACKEND_URL = "https://ssc-catalog-cap-backend.cfapps.us10-001.hana.ondemand.com"

def strip_html(text):
    """Remove HTML tags and decode entities, return plain text."""
    if not text:
        return ""
    if not isinstance(text, str):
        return str(text)
    # Remove HTML tags
    text = re.sub(r'<[^>]+>', ' ', text)
    # Decode common HTML entities
    text = text.replace('&nbsp;', ' ')
    text = text.replace('&amp;', '&')
    text = text.replace('&lt;', '<')
    text = text.replace('&gt;', '>')
    text = text.replace('&quot;', '"')
    text = text.replace('&#39;', "'")
    text = text.replace('&apos;', "'")
    # Collapse whitespace
    text = re.sub(r'\s+', ' ', text).strip()
    return text

def list_to_plain(value):
    """Convert list or string value to plain text."""
    if not value:
        return ""
    if isinstance(value, list):
        return "; ".join([strip_html(str(v)) for v in value if v])
    return strip_html(str(value))

def get_classification(service, feature_key):
    """Get list of values for a classificationFeatures key."""
    features = service.get("classificationFeatures", {})
    if not features:
        return []
    # classificationFeatures can be a dict or a list of {code, values} objects
    if isinstance(features, list):
        for item in features:
            if isinstance(item, dict) and item.get("code") == feature_key:
                vals = item.get("values", item.get("value", []))
                if isinstance(vals, list):
                    return [v.get("code", v) if isinstance(v, dict) else v for v in vals]
                return [vals] if vals else []
        return []
    # dict format
    val = features.get(feature_key, [])
    if isinstance(val, list):
        return val
    if val:
        return [val]
    return []

def get_supercategory_codes(service):
    """Get list of supercategory codes."""
    supercats = service.get("supercategories", [])
    if not supercats:
        return []
    codes = []
    for sc in supercats:
        code = sc.get("code", "") if isinstance(sc, dict) else str(sc)
        if code:
            codes.append(code)
    return codes

def main():
    print("Fetching snapshot from backend...")
    try:
        resp = requests.get(f"{BACKEND_URL}/api/catalog/getSnapshot", timeout=60)
        resp.raise_for_status()
        snapshot = resp.json()
    except Exception as e:
        print(f"❌ Failed to fetch snapshot: {e}")
        sys.exit(1)

    # payload is a double-encoded JSON string
    payload_raw = snapshot.get("payload", "")
    if isinstance(payload_raw, str):
        payload = json.loads(payload_raw)
    else:
        payload = payload_raw

    flat_index = payload.get("flat_index", payload.get("flatIndex", {}))
    if not flat_index:
        print("❌ No services in snapshot")
        sys.exit(1)

    services = list(flat_index.values())
    print(f"✅ Loaded {len(services)} services")

    # Collect all unique values for dynamic columns
    all_engagement_types = set()
    all_project_phases = set()
    all_supercategory_codes = set()

    for svc in services:
        for et in get_classification(svc, "engagementType"):
            all_engagement_types.add(et)
        for ph in get_classification(svc, "sapActivateProjectPhase"):
            all_project_phases.add(ph)
        for code in get_supercategory_codes(svc):
            all_supercategory_codes.add(code)

    engagement_types = sorted(all_engagement_types)
    project_phases = sorted(all_project_phases)
    supercategory_codes = sorted(all_supercategory_codes)

    print(f"  Engagement types: {len(engagement_types)}")
    print(f"  Project phases: {len(project_phases)}")
    print(f"  Supercategory codes: {len(supercategory_codes)}")

    # Build column headers
    fixed_columns = [
        "serviceCode",
        "serviceNumber",
        "serviceName",
        "serviceShortName",
        "crmBaseCategory",
        "summary",
        "serviceTeaserText",
        "businessNeeds",
        "keyBenefits",
        "deliveryApproach",
        "description",
        "scKeywords",
        "efforts",
    ]

    # Dynamic engagement type columns
    et_columns = [f"ET: {et}" for et in engagement_types]
    # Dynamic project phase columns
    ph_columns = [f"Phase: {ph}" for ph in project_phases]
    # Dynamic supercategory columns
    sc_columns = [f"SC: {code}" for code in supercategory_codes]

    all_columns = fixed_columns + et_columns + ph_columns + sc_columns

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Services"

    # Header style
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0070F2", end_color="0070F2", fill_type="solid")
    et_fill = PatternFill(start_color="107E3E", end_color="107E3E", fill_type="solid")
    ph_fill = PatternFill(start_color="E76500", end_color="E76500", fill_type="solid")
    sc_fill = PatternFill(start_color="5A1846", end_color="5A1846", fill_type="solid")

    # Write headers
    for col_idx, col_name in enumerate(all_columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        if col_name.startswith("ET:"):
            cell.fill = et_fill
        elif col_name.startswith("Phase:"):
            cell.fill = ph_fill
        elif col_name.startswith("SC:"):
            cell.fill = sc_fill
        else:
            cell.fill = header_fill

    # Write data rows
    for row_idx, svc in enumerate(services, 2):
        features = svc.get("classificationFeatures", {}) or {}
        svc_et = get_classification(svc, "engagementType")
        svc_ph = get_classification(svc, "sapActivateProjectPhase")
        svc_sc = get_supercategory_codes(svc)

        # crmBaseCategory
        crm_cats = get_classification(svc, "crmBaseCategory")
        crm_base = "; ".join(crm_cats) if crm_cats else ""

        # scKeywords
        sc_keywords = svc.get("scKeywords", [])
        if isinstance(sc_keywords, list):
            keywords_str = "; ".join(sc_keywords)
        else:
            keywords_str = str(sc_keywords) if sc_keywords else ""

        row_data = {
            "serviceCode": svc.get("code", "") or svc.get("serviceCode", ""),
            "serviceNumber": svc.get("serviceNumber", ""),
            "serviceName": svc.get("name", "") or svc.get("serviceName", ""),
            "serviceShortName": svc.get("serviceShortName", "") or svc.get("shortName", ""),
            "crmBaseCategory": crm_base,
            "summary": strip_html(svc.get("summary", "") or svc.get("shortDescription", "")),
            "serviceTeaserText": strip_html(svc.get("serviceTeaserText", "") or svc.get("teaserText", "")),
            "businessNeeds": strip_html(list_to_plain(svc.get("businessNeeds", "") or svc.get("businessNeed", ""))),
            "keyBenefits": strip_html(list_to_plain(svc.get("keyBenefits", "") or svc.get("benefits", ""))),
            "deliveryApproach": strip_html(list_to_plain(svc.get("deliveryApproach", "") or svc.get("approach", ""))),
            "description": strip_html(list_to_plain(svc.get("description", ""))),
            "scKeywords": keywords_str,
        "efforts": "; ".join(get_classification(svc, "effortEstimateDays")),
        }

        for col_idx, col_name in enumerate(all_columns, 1):
            if col_name in row_data:
                val = row_data[col_name]
            elif col_name.startswith("ET: "):
                et = col_name[4:]
                val = "v" if et in svc_et else ""
            elif col_name.startswith("Phase: "):
                ph = col_name[7:]
                val = "v" if ph in svc_ph else ""
            elif col_name.startswith("SC: "):
                sc = col_name[4:]
                val = "v" if sc in svc_sc else ""
            else:
                val = ""

            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if val == "v":
                cell.alignment = Alignment(horizontal="center")

    # Set column widths
    for col_idx, col_name in enumerate(all_columns, 1):
        col_letter = get_column_letter(col_idx)
        if col_name in ("summary", "businessNeeds", "keyBenefits", "deliveryApproach", "description", "serviceTeaserText"):
            ws.column_dimensions[col_letter].width = 50
        elif col_name in ("serviceName",):
            ws.column_dimensions[col_letter].width = 40
        elif col_name.startswith("ET:") or col_name.startswith("Phase:") or col_name.startswith("SC:"):
            ws.column_dimensions[col_letter].width = 20
        else:
            ws.column_dimensions[col_letter].width = 20

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(all_columns))}1"

    output_file = "ssc_services_export.xlsx"
    wb.save(output_file)
    print(f"\n✅ Excel exported: {output_file}")
    print(f"   {len(services)} services | {len(all_columns)} columns")
    print(f"   ({len(et_columns)} engagement type cols | {len(ph_columns)} project phase cols | {len(sc_columns)} supercategory cols)")

if __name__ == "__main__":
    main()
